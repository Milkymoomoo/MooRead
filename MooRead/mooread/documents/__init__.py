from __future__ import annotations

import hashlib
import io
import os
import re
import zipfile
from dataclasses import dataclass, field
from html import unescape
from pathlib import Path
from typing import Iterable, Optional
from xml.etree import ElementTree as ET

from pypdf import PdfReader
from striprtf.striprtf import rtf_to_text

from mooread.config import OCR_CACHE, TARGET_CHUNK_CHARS


IMAGE_EXTS = {".png", ".jpg", ".jpeg", ".webp", ".bmp", ".tif", ".tiff", ".gif"}
TEXT_EXTS = {".txt", ".md", ".html", ".htm", ".csv"}
SUPPORTED_EXTS = IMAGE_EXTS | TEXT_EXTS | {".pdf", ".epub", ".rtf", ".docx"}


@dataclass
class TocEntry:
    title: str
    href: str = ""
    position: int = 0


@dataclass
class DocumentBlock:
    """A display/speech unit. Audio is never stored here."""

    index: int
    text: str
    kind: str = "text"  # text | heading | image | page
    source: str = ""
    page: Optional[int] = None


@dataclass
class LoadedDocument:
    path: str
    title: str
    kind: str
    text: str
    blocks: list[DocumentBlock] = field(default_factory=list)
    toc: list[TocEntry] = field(default_factory=list)
    warnings: list[str] = field(default_factory=list)
    page_count: int = 1
    ocr_used: bool = False

    @property
    def char_count(self) -> int:
        return len(self.text)


_TAG_RE = re.compile(r"<[^>]+>")
_WS_RE = re.compile(r"[ \t]+")
_MULTI_NL = re.compile(r"\n{3,}")


def _clean_html(raw: str) -> str:
    raw = re.sub(r"(?is)<(script|style).*?>.*?</\1>", " ", raw)
    raw = re.sub(r"(?i)<br\s*/?>", "\n", raw)
    raw = re.sub(r"(?i)</(p|div|h[1-6]|li|tr|section|article)>", "\n\n", raw)
    raw = _TAG_RE.sub(" ", raw)
    raw = unescape(raw)
    raw = _WS_RE.sub(" ", raw)
    raw = _MULTI_NL.sub("\n\n", raw)
    return raw.strip()


def _file_key(path: Path) -> str:
    st = path.stat()
    blob = f"{path.resolve()}|{st.st_size}|{int(st.st_mtime)}"
    return hashlib.sha1(blob.encode("utf-8")).hexdigest()[:16]


def load_document(path: str | Path, ocr_engine=None) -> LoadedDocument:
    path = Path(path)
    if not path.exists():
        raise FileNotFoundError(path)
    ext = path.suffix.lower()
    if ext == ".epub":
        return _load_epub(path)
    if ext == ".pdf":
        return _load_pdf(path, ocr_engine)
    if ext == ".rtf":
        return _load_rtf(path)
    if ext == ".docx":
        return _load_docx(path)
    if ext in IMAGE_EXTS:
        return _load_image(path, ocr_engine)
    if ext in TEXT_EXTS or ext == "":
        return _load_text(path)
    raise ValueError(f"Unsupported file type: {ext}")


def document_from_text(text: str, title: str = "Camera", kind: str = "camera") -> LoadedDocument:
    body = (text or "").strip()
    blocks = split_display_blocks(body) if body else []
    if not blocks and body:
        blocks = [DocumentBlock(index=0, text=body, kind=kind)]
    return LoadedDocument(
        path="",
        title=title,
        kind=kind,
        text=body,
        blocks=blocks,
        ocr_used=kind in {"camera", "image", "manga"},
    )


def _load_text(path: Path) -> LoadedDocument:
    data = path.read_bytes()
    text = None
    for enc in ("utf-8-sig", "utf-8", "cp1252", "latin-1"):
        try:
            text = data.decode(enc)
            break
        except UnicodeDecodeError:
            continue
    if text is None:
        text = data.decode("utf-8", errors="replace")
    if path.suffix.lower() in {".html", ".htm"}:
        text = _clean_html(text)
    blocks = split_display_blocks(text)
    return LoadedDocument(
        path=str(path),
        title=path.stem,
        kind="text",
        text=text,
        blocks=blocks,
        page_count=max(1, (len(text) // 1800) + 1),
    )


def _load_rtf(path: Path) -> LoadedDocument:
    raw = path.read_text(encoding="latin-1", errors="ignore")
    text = rtf_to_text(raw)
    text = _MULTI_NL.sub("\n\n", text).strip()
    return LoadedDocument(
        path=str(path),
        title=path.stem,
        kind="rtf",
        text=text,
        blocks=split_display_blocks(text),
        page_count=max(1, (len(text) // 1800) + 1),
    )


def _load_docx(path: Path) -> LoadedDocument:
    try:
        from docx import Document
    except ImportError as exc:
        raise RuntimeError("python-docx is required for .docx") from exc
    doc = Document(str(path))
    parts: list[str] = []
    blocks: list[DocumentBlock] = []
    i = 0
    for para in doc.paragraphs:
        t = para.text.strip()
        if not t:
            continue
        style = (para.style.name or "").lower() if para.style else ""
        kind = "heading" if "heading" in style else "text"
        blocks.append(DocumentBlock(index=i, text=t, kind=kind, source=str(path)))
        parts.append(t)
        i += 1
    text = "\n\n".join(parts)
    return LoadedDocument(
        path=str(path),
        title=path.stem,
        kind="docx",
        text=text,
        blocks=blocks or split_display_blocks(text),
    )


def _load_image(path: Path, ocr_engine=None) -> LoadedDocument:
    text = ""
    ocr_used = False
    warnings: list[str] = []
    if ocr_engine is not None:
        text = ocr_engine.recognize_file(path)
        ocr_used = bool(text.strip())
    if not text.strip():
        warnings.append("No text detected in image. OCR engine may be missing or the image has no readable text.")
        text = f"[Image: {path.name}]"
    blocks = [DocumentBlock(index=0, text=text, kind="image", source=str(path))]
    return LoadedDocument(
        path=str(path),
        title=path.stem,
        kind="image",
        text=text,
        blocks=blocks,
        warnings=warnings,
        ocr_used=ocr_used,
    )


def _load_pdf(path: Path, ocr_engine=None) -> LoadedDocument:
    reader = PdfReader(str(path))
    page_texts: list[str] = []
    blocks: list[DocumentBlock] = []
    ocr_used = False
    warnings: list[str] = []
    for i, page in enumerate(reader.pages):
        extracted = (page.extract_text() or "").strip()
        kind = "page"
        if len(extracted) < 40:
            # Defer OCR until the voice cursor reaches this page.
            kind = "page-scan"
        if not extracted:
            extracted = ""
        page_texts.append(extracted)
        if extracted or kind == "page-scan":
            blocks.append(DocumentBlock(index=len(blocks), text=extracted, kind=kind, source=str(path), page=i + 1))
    text = "\n\n".join(page_texts)
    title = path.stem
    try:
        meta = reader.metadata
        if meta and meta.title:
            title = str(meta.title)
    except Exception:
        pass
    return LoadedDocument(
        path=str(path),
        title=title,
        kind="pdf",
        text=text,
        blocks=blocks,
        toc=[TocEntry(title=f"Page {i+1}", position=i) for i in range(len(blocks))],
        warnings=warnings,
        page_count=len(reader.pages),
        ocr_used=ocr_used,
    )


def _load_epub(path: Path) -> LoadedDocument:
    from ebooklib import epub, ITEM_DOCUMENT

    book = epub.read_epub(str(path))
    title = path.stem
    try:
        titles = book.get_metadata("DC", "title")
        if titles:
            title = titles[0][0]
    except Exception:
        pass

    spine_ids = []
    for item in book.spine:
        if isinstance(item, tuple):
            spine_ids.append(item[0])
        else:
            spine_ids.append(item)

    id_map = {item.get_id(): item for item in book.get_items()}
    blocks: list[DocumentBlock] = []
    toc: list[TocEntry] = []
    parts: list[str] = []
    idx = 0
    for sid in spine_ids:
        item = id_map.get(sid)
        if item is None:
            continue
        try:
            html = item.get_content().decode("utf-8", errors="ignore")
        except Exception:
            continue
        text = _clean_html(html)
        if not text:
            continue
        heading = text.split("\n", 1)[0][:80]
        toc.append(TocEntry(title=heading, href=sid, position=idx))
        blocks.append(DocumentBlock(index=idx, text=text, kind="text", source=sid))
        parts.append(text)
        idx += 1

    if not parts:
        # fallback: any HTML documents not in spine
        for item in book.get_items_of_type(ITEM_DOCUMENT):
            html = item.get_content().decode("utf-8", errors="ignore")
            text = _clean_html(html)
            if text:
                parts.append(text)
                blocks.append(DocumentBlock(index=len(blocks), text=text, kind="text"))

    text = "\n\n".join(parts)
    return LoadedDocument(
        path=str(path),
        title=title,
        kind="epub",
        text=text,
        blocks=blocks or split_display_blocks(text),
        toc=toc,
        page_count=max(1, len(blocks)),
    )


def split_display_blocks(text: str) -> list[DocumentBlock]:
    paras = [p.strip() for p in re.split(r"\n\s*\n", text) if p.strip()]
    blocks = []
    for i, p in enumerate(paras):
        kind = "heading" if (len(p) < 80 and not p.endswith((".", "!", "?", '"', "'"))) else "text"
        blocks.append(DocumentBlock(index=i, text=p, kind=kind))
    if not blocks and text.strip():
        blocks.append(DocumentBlock(index=0, text=text.strip()))
    return blocks


def iter_speech_chunks(text: str, target: int = TARGET_CHUNK_CHARS) -> Iterable[str]:
    """Split into speakable blocks that respect sentences and paragraphs."""
    from mooread.config import MAX_CHUNK_CHARS, MIN_CHUNK_CHARS

    paragraphs = [p.strip() for p in re.split(r"\n\s*\n", text) if p.strip()]
    buf: list[str] = []
    size = 0
    for para in paragraphs:
        sentences = _split_sentences(para)
        for sent in sentences:
            if size and size + len(sent) + 1 > target and size >= MIN_CHUNK_CHARS:
                yield " ".join(buf).strip()
                buf, size = [], 0
            if len(sent) > MAX_CHUNK_CHARS:
                if buf:
                    yield " ".join(buf).strip()
                    buf, size = [], 0
                for piece in _hard_wrap(sent, MAX_CHUNK_CHARS):
                    yield piece
                continue
            buf.append(sent)
            size += len(sent) + 1
        if size >= target:
            yield " ".join(buf).strip()
            buf, size = [], 0
    if buf:
        yield " ".join(buf).strip()


def _split_sentences(text: str) -> list[str]:
    parts = re.split(r"(?<=[.!?])\s+(?=[A-Z\"“‘])", text)
    out: list[str] = []
    for p in parts:
        p = p.strip()
        if p:
            out.append(p)
    return out or ([text.strip()] if text.strip() else [])


def _hard_wrap(text: str, limit: int) -> list[str]:
    words = text.split()
    if len(words) <= 1 and len(text) > limit:
        return [text[i : i + limit] for i in range(0, len(text), limit)]
    rows: list[str] = []
    cur: list[str] = []
    n = 0
    for w in words:
        if n + len(w) + 1 > limit and cur:
            rows.append(" ".join(cur))
            cur, n = [], 0
        cur.append(w)
        n += len(w) + 1
    if cur:
        rows.append(" ".join(cur))
    return rows
